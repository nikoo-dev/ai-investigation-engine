"""Small, format-specific parsers with source locations and explicit failures."""

from __future__ import annotations

import io
import json
import re
import unicodedata
import zipfile
from pathlib import PureWindowsPath
from typing import Any

from docx import Document as WordDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from investigation.errors import InvalidDocument
from investigation.models import ParsedDocument, SourceLocation, TextBlock

SUPPORTED_FORMATS = {".pdf", ".docx", ".txt", ".md", ".jsonl"}
MAX_EXTRACTED_CHARS = 10_000_000
MAX_ARCHIVE_BYTES = 50_000_000


def safe_filename(filename: str) -> str:
    name = PureWindowsPath(filename.replace("/", "\\")).name
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip(" .")[:240]
    if not name:
        raise InvalidDocument("A filename with a supported extension is required")
    return name


def normalize_text(text: str) -> str:
    """Keep original Unicode text readable while removing extraction noise."""
    text = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    return text.replace("\x00", "").strip()


def _assemble(
    title: str, pieces: list[tuple[str, SourceLocation]], metadata: dict[str, Any]
) -> ParsedDocument:
    blocks: list[TextBlock] = []
    texts: list[str] = []
    position = 0
    for raw, location in pieces:
        text = normalize_text(raw)
        if not text:
            continue
        if texts:
            position += 2
        location = location.model_copy(
            update={"start_char": position, "end_char": position + len(text)}
        )
        blocks.append(TextBlock(text=text, location=location))
        texts.append(text)
        position += len(text)
        if position > MAX_EXTRACTED_CHARS:
            raise InvalidDocument("Extracted text exceeds the supported document size")
    if not texts:
        raise InvalidDocument("No readable text was found. Scanned PDFs require OCR before upload.")
    return ParsedDocument(title=title, text="\n\n".join(texts), blocks=blocks, metadata=metadata)


def parse_document(
    filename: str, content: bytes, max_bytes: int = 10 * 1024 * 1024
) -> ParsedDocument:
    filename = safe_filename(filename)
    suffix = PureWindowsPath(filename).suffix.lower()
    if suffix not in SUPPORTED_FORMATS:
        raise InvalidDocument("Supported formats are PDF, DOCX, TXT, Markdown and JSONL")
    if not content:
        raise InvalidDocument("The uploaded file is empty")
    if len(content) > max_bytes:
        raise InvalidDocument(f"File exceeds the {max_bytes // (1024 * 1024)} MB upload limit")
    title = PureWindowsPath(filename).stem
    if suffix == ".pdf":
        return _parse_pdf(title, content)
    if suffix == ".docx":
        return _parse_docx(title, content)
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise InvalidDocument("Text and JSONL files must use UTF-8 encoding") from exc
    if suffix == ".jsonl":
        pieces: list[tuple[str, SourceLocation]] = []
        try:
            for index, line in enumerate(text.splitlines(), 1):
                if not line.strip():
                    continue
                row = json.loads(line)
                if not isinstance(row, dict) or not isinstance(row.get("text"), str):
                    raise ValueError("Each row must be an object containing a string text field")
                pieces.append(
                    (row["text"], SourceLocation(section=str(row.get("title", f"Record {index}"))))
                )
        except (ValueError, TypeError) as exc:
            raise InvalidDocument(f"Invalid JSONL: {exc}") from exc
        return _assemble(title, pieces, {"parser": "jsonl", "records": len(pieces)})
    pieces = []
    section: str | None = None
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = normalize_text(paragraph)
        if suffix == ".md" and paragraph.startswith("#"):
            section = paragraph.splitlines()[0].lstrip("# ").strip()
        pieces.append((paragraph, SourceLocation(section=section)))
    return _assemble(title, pieces, {"parser": "utf8", "format": suffix.lstrip(".")})


def _parse_pdf(title: str, content: bytes) -> ParsedDocument:
    if not content.lstrip().startswith(b"%PDF-"):
        raise InvalidDocument("The file extension is PDF but the content is not a PDF")
    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise InvalidDocument("Encrypted PDFs must be decrypted before upload")
        if len(reader.pages) > 1000:
            raise InvalidDocument("PDFs are limited to 1000 pages per upload")
        metadata: Any = reader.metadata or {}
        pieces = [
            (page.extract_text() or "", SourceLocation(page=index + 1))
            for index, page in enumerate(reader.pages)
        ]
        return _assemble(
            str(metadata.get("/Title") or title),
            pieces,
            {
                "parser": "pypdf",
                "pages": len(reader.pages),
                "author": str(metadata.get("/Author") or ""),
            },
        )
    except InvalidDocument:
        raise
    except Exception as exc:
        raise InvalidDocument(
            "The PDF could not be parsed; it may be damaged or unsupported"
        ) from exc


def _parse_docx(title: str, content: bytes) -> ParsedDocument:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 10000
                or sum(entry.file_size for entry in entries) > MAX_ARCHIVE_BYTES
            ):
                raise InvalidDocument("DOCX archive expands beyond the supported size")
            if any(
                entry.file_size > 1_000_000 and entry.file_size / max(entry.compress_size, 1) > 1000
                for entry in entries
            ):
                raise InvalidDocument("DOCX archive has an unsafe compression ratio")
            if "word/document.xml" not in archive.namelist():
                raise InvalidDocument("The file is not a valid Word DOCX document")
        word = WordDocument(io.BytesIO(content))
        pieces: list[tuple[str, SourceLocation]] = []
        section: str | None = None
        for child in word.element.body.iterchildren():
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, word)
                style = paragraph.style.name if paragraph.style else ""
                if style.startswith("Heading") or style == "Title":
                    section = paragraph.text.strip()
                pieces.append((paragraph.text, SourceLocation(section=section)))
            elif isinstance(child, CT_Tbl):
                table = Table(child, word)
                for row in table.rows:
                    pieces.append(
                        (
                            " | ".join(cell.text for cell in row.cells),
                            SourceLocation(section=section),
                        )
                    )
        props = word.core_properties
        return _assemble(
            props.title or title,
            pieces,
            {"parser": "python-docx", "author": props.author or "", "subject": props.subject or ""},
        )
    except InvalidDocument:
        raise
    except Exception as exc:
        raise InvalidDocument(
            "The DOCX could not be parsed; it may be damaged or unsupported"
        ) from exc
