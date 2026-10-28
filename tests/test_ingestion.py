import io
from uuid import uuid4

import pytest
from docx import Document as WordDocument

from investigation.config import Settings
from investigation.embeddings import HashEmbedder
from investigation.errors import ConflictError, InvalidDocument, NotFound
from investigation.ingestion import Ingestor
from investigation.models import SearchFilters
from investigation.parsers import parse_document, safe_filename
from investigation.storage import Store


@pytest.fixture
def ingest(tmp_path):
    settings = Settings(
        _env_file=None,
        environment="test",
        data_dir=tmp_path,
        embedding_backend="hash",
        embedding_dimension=64,
        chunk_tokens=32,
        chunk_overlap=4,
        reranker_enabled=False,
    )
    embedder = HashEmbedder(dimension=64)
    store = Store(settings, embedder.dimension)
    yield Ingestor(store, embedder, settings)
    store.close()


def test_exact_duplicates_do_not_inflate_count(ingest):
    first = ingest.ingest("one.txt", b"Mars is a planet.")
    second = ingest.ingest("two.txt", b"  MARS  is a planet.\n")
    assert second.status == "duplicate"
    assert second.document.id == first.document.id
    assert ingest.store.stats()["documents"] == 1


def test_replacement_rejects_missing_document(ingest):
    with pytest.raises(NotFound):
        ingest.ingest("one.txt", b"A statement.", document_id=str(uuid4()))


def test_docx_paragraph_and_table_locations():
    word = WordDocument()
    word.add_heading("Incident summary", 1)
    word.add_paragraph("Payments failed on March 2.")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Payment gateway"
    content = io.BytesIO()
    word.save(content)
    parsed = parse_document("incident.docx", content.getvalue())
    assert "Service | Payment gateway" in parsed.text
    assert parsed.blocks[1].location.section == "Incident summary"
    for block in parsed.blocks:
        assert block.text == parsed.text[block.location.start_char : block.location.end_char]


@pytest.mark.parametrize(
    "filename,content",
    [
        ("bad.pdf", b"not pdf"),
        ("bad.docx", b"not a zip"),
        ("empty.txt", b""),
        ("unknown.exe", b"123"),
        ("text.txt", b"\xff\xfe"),
        ("bad.jsonl", b'{"text":5}'),
    ],
)
def test_invalid_files_are_explicit_errors(filename, content):
    with pytest.raises(InvalidDocument):
        parse_document(filename, content)


def test_upload_size_limit():
    with pytest.raises(InvalidDocument, match="upload limit"):
        parse_document("text.txt", b"x" * 100, max_bytes=10)


def test_unicode_and_filename_sanitization():
    assert safe_filename("../../danger.txt") == "danger.txt"
    assert safe_filename("C:\\upload\\incident.txt") == "incident.txt"
    text = "ქართული ტექსტი\n\nSecond paragraph."
    parsed = parse_document("unicode.txt", text.encode("utf-8"))
    assert parsed.text == text
    assert parsed.blocks[1].location.start_char == text.index("Second")


def test_batch_invalid_rows_and_duplicates_are_counted(ingest):
    counts = ingest.ingest_records(
        [
            {"id": str(uuid4()), "text": "A fact."},
            {"id": str(uuid4()), "text": "A fact."},
            {"text": ""},
        ]
    )
    assert counts["processed"] == 3
    assert counts["created"] == 1
    assert counts["duplicates"] == 1
    assert counts["failed"] == 1


def test_batch_progress_is_checkpointed_and_callback_failure_is_not_swallowed(ingest):
    records = [{"id": str(uuid4()), "text": f"Source record number {index}."} for index in range(5)]
    snapshots = []
    result = ingest.ingest_records(records, batch_size=2, progress=snapshots.append)
    assert [snapshot["processed"] for snapshot in snapshots] == [2, 4, 5]
    assert snapshots[-1] == result

    def fail_progress(_):
        raise RuntimeError("checkpoint unavailable")

    with pytest.raises(RuntimeError, match="checkpoint unavailable"):
        ingest.ingest_records([{"text": "A new persisted source."}], progress=fail_progress)
    assert ingest.store.stats()["documents"] == 6


def test_ingestion_reserves_title_capacity_before_model_token_chunking(ingest):
    class CharacterEmbedder(HashEmbedder):
        max_content_tokens = 48

        @staticmethod
        def token_spans(value):
            return [(index, index + 1) for index, char in enumerate(value) if not char.isspace()]

        def embed(self, texts):
            assert all(len(self.token_spans(text)) <= self.max_content_tokens for text in texts)
            return super().embed(texts)

    ingest.embedder = CharacterEmbedder(dimension=64)
    title = "A deliberately extremely long document title which cannot consume the entire context"
    text = "antidisestablishmentarianism repeated verylongunknownword tailfact"
    result = ingest.ingest_record({"title": title, "text": text})
    chunks = ingest.store.query_dense(
        ingest.embedder.embed(["tailfact"])[0],
        20,
        SearchFilters(),
    )
    assert result.document.title == title
    assert len(chunks) > 1
    assert any(hit.chunk.location.end_char == len(text) for hit in chunks)
    assert all(
        len(ingest.embedder.token_spans(hit.chunk.metadata["embedding_title"])) <= 12
        for hit in chunks
    )


def test_incomplete_sentence_labels_do_not_discard_source_text(ingest):
    text = "The unlabeled first fact. The labeled second fact. A third important fact."
    result = ingest.ingest_record(
        {
            "text": text,
            "metadata": {"sentences": ["The labeled second fact.", "Missing fake sentence."]},
        }
    )
    chunks = ingest.store.query_dense(ingest.embedder.embed(["fact"])[0], 20, SearchFilters())
    covered = {
        index
        for hit in chunks
        for index in range(hit.chunk.location.start_char, hit.chunk.location.end_char)
    }
    assert all(
        index in covered for index, char in enumerate(result.document.text) if not char.isspace()
    )


def test_same_batch_repeated_identity_is_duplicate_only_if_content_matches(ingest):
    source_id = str(uuid4())
    first = {"id": source_id, "text": "The same source content."}
    counts = ingest.ingest_records(
        [
            first,
            first,
            {"id": source_id, "text": "A conflicting different content for the same identity."},
        ]
    )
    assert counts["created"] == 1
    assert counts["duplicates"] == 1
    assert counts["failed"] == 1


def test_delete_during_embedding_does_not_resurrect_replacement(ingest, monkeypatch):
    created = ingest.ingest("original.txt", b"The original document is active.")
    original_vectors = ingest._vectors

    def delete_then_embed(chunks):
        assert ingest.store.delete_document(created.document.id)
        return original_vectors(chunks)

    monkeypatch.setattr(ingest, "_vectors", delete_then_embed)
    with pytest.raises(ConflictError) as error:
        ingest.ingest(
            "replacement.txt",
            b"A late replacement must remain unpublished.",
            document_id=created.document.id,
        )
    assert error.value.status_code == 409
    assert ingest.store.get_document(created.document.id) is None
    assert ingest.store.stats()["documents"] == 0
    assert ingest.store.stats()["chunks"] == 0


def test_competing_revision_wins_and_stale_cleanup_does_not_remove_it(ingest, monkeypatch):
    created = ingest.ingest("original.txt", b"The original document is active.")
    original_vectors = ingest._vectors
    competing = Ingestor(ingest.store, ingest.embedder, ingest.settings)
    winner = []

    def publish_competing_then_embed(chunks):
        winner.append(
            competing.ingest(
                "winner.txt",
                b"The competing revision was published first.",
                document_id=created.document.id,
            )
        )
        return original_vectors(chunks)

    monkeypatch.setattr(ingest, "_vectors", publish_competing_then_embed)
    with pytest.raises(ConflictError):
        ingest.ingest(
            "stale.txt",
            b"A stale revision must not overwrite the competing update.",
            document_id=created.document.id,
        )
    active = ingest.store.get_document(created.document.id)
    assert active.revision == 2
    assert active.text == winner[0].document.text
    hits = ingest.store.query_dense(
        ingest.embedder.embed(["competing revision"])[0], 10, SearchFilters()
    )
    assert hits
    assert all(hit.chunk.revision == 2 and "competing revision" in hit.chunk.text for hit in hits)
    assert ingest.store.stats()["chunks"] == winner[0].document.chunk_count


def test_batch_delete_between_preparation_and_commit_is_conflict(ingest, monkeypatch):
    source_id = str(uuid4())
    created = ingest.ingest_record({"id": source_id, "text": "The existing dataset source."})
    original_vectors = ingest._vectors

    def delete_then_embed(chunks):
        assert ingest.store.delete_document(created.document.id)
        return original_vectors(chunks)

    monkeypatch.setattr(ingest, "_vectors", delete_then_embed)
    with pytest.raises(ConflictError):
        ingest.ingest_records(
            [{"id": source_id, "text": "The late dataset update cannot restore a deleted source."}]
        )
    assert ingest.store.stats()["documents"] == 0
    assert ingest.store.stats()["chunks"] == 0


def test_dataset_replay_does_not_restore_deleted_identity(ingest):
    record = {"id": str(uuid4()), "text": "A deliberately deleted source document."}
    created = ingest.ingest_record(record)
    assert ingest.store.delete_document(created.document.id)
    with pytest.raises(ConflictError, match="deleted"):
        ingest.ingest_record(record)
    counts = ingest.ingest_records([record])
    assert counts["failed"] == 1
    assert counts["created"] == 0
    assert ingest.store.stats()["documents"] == 0


def test_new_identity_cannot_overwrite_a_tombstone(ingest):
    result = ingest.ingest("original.txt", b"Original source text.")
    assert ingest.store.delete_document(result.document.id)
    late = result.document.model_copy(update={"text": "Unexpected resurrection."})
    chunks = ingest._chunks(late)
    dense, sparse = ingest._vectors(chunks)
    with pytest.raises(ConflictError):
        ingest.store.commit_document(late, chunks, dense, sparse)
    assert ingest.store.stats()["documents"] == 0
    assert ingest.store.stats()["chunks"] == 0
